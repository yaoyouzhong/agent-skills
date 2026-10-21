"""Write an audited XLSX plan using openpyxl; reject unsupported native objects."""
import copy,warnings
from pathlib import Path
from zipfile import ZipFile


def write_plan(plan):
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.formula.translate import Translator
    from openpyxl.utils.cell import range_boundaries
    source=Path(plan['template']);out=Path(plan['output'])
    if out.exists():raise ValueError('Use a new Excel output path.')
    blocked=('xl/drawings/','xl/media/','xl/charts/','xl/pivot','xl/embeddings/','xl/activeX/','xl/vba','xl/externalLinks/','xl/slicer','customXml/')
    with ZipFile(source) as z:
        if any(n.startswith(blocked) for n in z.namelist()):raise ValueError('Portable XLSX backend requires a cell-only template; native objects require a task-specific adapter.')
        if any(b'<extLst' in z.read(n) or b':extLst' in z.read(n) for n in z.namelist() if n.startswith('xl/') and n.endswith('.xml')):raise ValueError('Unsupported XLSX extensions require a task-specific adapter.')
    with warnings.catch_warnings():
        warnings.simplefilter('error')
        wb=openpyxl.load_workbook(source)
    ws=wb[plan['sheet']]
    if ws.protection.sheet:raise ValueError('Protected target worksheet requires an editable template.')
    snapshots={}
    for item in plan['copies']:
        address=item['from']
        if address in snapshots:continue
        left,top,right,bottom=range_boundaries(address)
        cells=[]
        for row in ws.iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right):
            for c in row:cells.append((c.row-top,c.column-left,c.value,copy.copy(c._style),copy.copy(c.alignment),copy.copy(c.hyperlink) if not isinstance(c,MergedCell) else None,copy.copy(c.comment),c.coordinate))
        merges=[(m.min_col-left,m.min_row-top,m.max_col-left,m.max_row-top) for m in ws.merged_cells.ranges if m.min_col>=left and m.max_col<=right and m.min_row>=top and m.max_row<=bottom]
        heights=[copy.copy(ws.row_dimensions[i]) for i in range(top,bottom+1)]
        snapshots[address]=(left,top,right,bottom,cells,merges,heights)
    for item in plan['copies']:
        left,top,right,bottom,cells,merges,heights=snapshots[item['from']]
        dl,dt,dr,db=range_boundaries(item['to'])
        if (right-left,bottom-top)!=(dr-dl,db-dt):raise ValueError('Copy rectangle dimensions differ.')
        for m in list(ws.merged_cells.ranges):
            if m.max_col<dl or m.min_col>dr or m.max_row<dt or m.min_row>db:continue
            if m.min_col<dl or m.max_col>dr or m.min_row<dt or m.max_row>db:raise ValueError('Copy intersects only part of a merged region.')
            ws.unmerge_cells(str(m))
        for ro,co,value,style,alignment,hyperlink,comment,original in cells:
            c=ws.cell(dt+ro,dl+co)
            c.value=Translator(value,origin=original).translate_formula(c.coordinate) if isinstance(value,str) and value.startswith('=') else value
            c._style=copy.copy(style);c.alignment=copy.copy(alignment)
            if hyperlink:hyperlink=copy.copy(hyperlink);hyperlink.ref=c.coordinate
            c._hyperlink=hyperlink;c.comment=copy.copy(comment)
        for i,h in enumerate(heights):h=copy.copy(h);h.index=dt+i;ws.row_dimensions[dt+i]=h
        for l,t,r,b in merges:ws.merge_cells(start_row=dt+t,start_column=dl+l,end_row=dt+b,end_column=dl+r)
    for address in plan['clearRanges']:
        for row in ws[address]:
            for c in row:
                if not isinstance(c,MergedCell):c.value=None
    for address in plan['unmergeRanges']:
        if address in ws.merged_cells:ws.unmerge_cells(address)
    for address in plan['merges']:ws.merge_cells(address)
    for item in plan['cells']:
        c=ws[item['cell']]
        if isinstance(c,MergedCell):raise ValueError('Target value is not at a merged region anchor: '+item['cell'])
        c.value=str(item['value']) if item['value'] is not None else ''
        c.data_type='s';c.number_format='@'
        alignment=copy.copy(c.alignment);alignment.wrap_text=True;c.alignment=alignment
    for item in plan['rowHeights']:ws.row_dimensions[item['row']].height=item['height']
    wb.save(out)
    return {'backend':'openpyxl','status':'exported_requires_native_render','output':str(out),'cells':len(plan['cells'])}
