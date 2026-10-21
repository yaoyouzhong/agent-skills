import json,sys,tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from template_prepare import clean_native
from template_docx import make_person,q
from lxml import etree as E

def run():
    checks={}
    with tempfile.TemporaryDirectory() as td:
        w=Path(td);doc=Document();t=doc.add_table(rows=3,cols=2)
        for row,values in zip(t.rows,[['姓名','示例甲'],['项目名称','旧项目'],['项目名称','旧项目二']]):
            for c,v in zip(row.cells,values):c.text=v
        native=w/'native.docx';doc.save(native)
        plan={'blank_cells':[{'row':0,'cell':1,'expected_text':'示例甲'},{'row':1,'cell':1,'expected_text':'旧项目'}],
              'remove_rows':[2],'keep_text':['姓名','项目名称']}
        out=w/'blank.docx';clean_native(native,plan,out)
        text=''.join(c.text for row in Document(out).tables[0].rows for c in row.cells)
        assert text=='姓名项目名称' and len(Document(out).tables[0].rows)==2
        checks['sample_values_removed_labels_and_repeat_prototype_retained']=True
        with ZipFile(native) as a,ZipFile(out) as b:
            assert a.namelist()==b.namelist() and all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')
        checks['unaffected_package_parts_are_identical']=True
        def reject(p,name):
            try:clean_native(native,p,w/name)
            except ValueError:return
            raise AssertionError('Unsafe template preparation accepted')
        reject({**plan,'blank_cells':plan['blank_cells'][:1]},'leak.docx')
        checks['remaining_personal_sample_text_rejected']=True
        reject({**plan,'blank_cells':[{'row':0,'cell':1,'expected_text':'other person'}]},'stale.docx')
        checks['changed_sample_value_rejected']=True
        with ZipFile(native) as z:body=E.fromstring(z.read('word/document.xml')).find(q('body'))
        row=body.find(q('tbl')).find(q('tr'));pr=E.SubElement(row,q('trPr'));h=E.SubElement(pr,q('trHeight'));h.set(q('val'),'2400')
        person={'id':'p','fields':{'skills':'短技能说明'}}
        mapping={'tables':[{'index':0,'fields':[], 'append_rows':[{'rows':[0,0],'auto_height':True,'fields':[{'row':0,'cell':0,'key':'skills'}]}]}]}
        result=make_person(body,mapping,person);table=next(n for n in result if n.tag==q('tbl'))
        assert table.findall(q('tr'))[0].find('w:trPr/w:trHeight',{'w':q('tr')[1:].split('}')[0]}) is not None
        assert table.findall(q('tr'))[-1].find('w:trPr/w:trHeight',{'w':q('tr')[1:].split('}')[0]}) is None
        assert '短技能说明' in ''.join(table.findall(q('tr'))[-1].itertext())
        checks['short_additional_field_can_grow_without_inheriting_tall_sample_height']=True
    print(json.dumps({'status':'passed','checks':checks}))

if __name__=='__main__':run()
