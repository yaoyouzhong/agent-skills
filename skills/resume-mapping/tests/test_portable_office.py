import json,sys,tempfile
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from openpyxl import Workbook
from lxml import etree as E
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from template_prepare import prepare
from extract_sources import sha
from libreoffice_bridge import select_sheet,find_executable
from environment_check import inspect


def run():
    checks={}
    with tempfile.TemporaryDirectory() as folder:
        w=Path(folder);doc=Document()
        for value in ('另一个人','选中样例'):
            table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='姓名';table.cell(0,1).text=value
        source=w/'source.docx';doc.save(source);before=sha(source)
        plan={'source_sha256':before,'table_index':1,'blank_cells':[{'row':0,'cell':1,'expected_text':'选中样例'}],'keep_text':['姓名']}
        report=prepare(source,plan,w/'prepared',backend='ooxml')
        final=Document(report['output']);assert len(final.tables)==1 and final.tables[0].cell(0,0).text=='姓名' and final.tables[0].cell(0,1).text==''
        assert sha(source)==before;checks['portable_sample_preparation_selects_correct_zero_based_table_and_preserves_original']=True
        doc.sections[0].header.paragraphs[0].text='其他样例姓名';leak=w/'leak.docx';doc.save(leak)
        try:prepare(leak,{**plan,'source_sha256':sha(leak)},w/'rejected',backend='ooxml')
        except ValueError:pass
        else:raise AssertionError('Unselected header content silently copied.')
        checks['sample_information_outside_selected_table_is_rejected']=True
        wb=Workbook();wb.active.title='其他人';wb.create_sheet('目标');xlsx=w/'source.xlsx';wb.save(xlsx);before=sha(xlsx)
        out=w/'selected.xlsx';select_sheet(xlsx,out,'目标')
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with ZipFile(xlsx) as a,ZipFile(out) as b:
            root=E.fromstring(b.read('xl/workbook.xml'));sheets=root.findall('s:sheets/s:sheet',ns)
            assert [(s.get('name'),s.get('state')) for s in sheets]==[('其他人','hidden'),('目标','visible')]
            assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='xl/workbook.xml')
        assert sha(xlsx)==before;checks['libreoffice_task_copy_selects_only_target_sheet_without_touching_other_parts_or_original']=True
        env=inspect();assert env['application_launch_performed'] is False and env['renderer_runtime_tested'] is False
        checks['environment_probe_does_not_claim_a_native_application_test']=True
        assert find_executable(w/'missing-soffice') is None or find_executable(w/'missing-soffice').is_file()
        checks['renderer_discovery_checks_existing_file']=True
    print(json.dumps({'status':'passed','checks':checks}))


if __name__=='__main__':run()
