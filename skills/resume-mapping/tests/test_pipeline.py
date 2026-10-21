"""Private-data-free behavioral regression. Run with a NEW --workdir."""
import argparse,copy,json,sys
from pathlib import Path
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from docx import Document
from apply_policies import apply
from extract_sources import sha
from source_audit import audit,dump,units_for
from record_evidence import locate
from template_docx import build,verify
from field_facts import classify
from review_ledger import create,verify as check_receipt
from record_checks import education_binding


def confirmed_fixture(data):
    """Only synthetic test data: its completion state is part of the test premise."""
    data['processing_date']='2026-10-20'
    for p in data['people']:
        p.setdefault('education_status_evidence',{})
        for i,e in enumerate(p['fields'].get('education',[])):
            if e.get('status')=='completed':
                p['education_status_evidence'][f'education[{i}]']={'basis':'user_confirmation','record_sha256':education_binding(e),
                    'reviewer':'synthetic_fixture','confirmation_reference':'private-data-free behavioral fixture',
                    'reason':'Synthetic test premise explicitly supplies completed status.'}
    return data


def denied(call):
    try:call()
    except (ValueError,KeyError,StopIteration):return True
    raise AssertionError('An invalid operation was accepted.')


def run(root):
    root.mkdir(parents=True,exist_ok=False);results={}
    source=root/'source.txt';source.write_text('测试人员\n示例大学\n本科\n2020\n项目甲\n为示例银行建设业务平台\n实现甲业务接口，保留成果100%\n项目乙\n为示例证券建设行情平台\n实现乙业务接口，保留成果90%\n',encoding='utf8')
    manifest={'sources':[{'id':'s1','source':str(source.resolve()),'sha256':sha(source),'text_file':str(source.resolve()),'status':'ok','detail_file':str((root/'detail.json').resolve())}]}
    dump(root/'detail.json',{})
    projects=[{'name':'项目甲','desc':'为示例银行建设业务平台','client':'','resp':'实现甲业务接口，保留成果100%'},{'name':'项目乙','desc':'为示例证券建设行情平台','client':'','resp':'实现乙业务接口，保留成果90%'}]
    data={'people':[{'id':'p1','source_ids':['s1'],'fields':{'meta':{'name':'测试人员'},'education':[{'school':'示例大学','level':'本科','end':'2020','degree':'','status':'completed'}],'projects':projects}}]}
    data,_,_=apply(confirmed_fixture(data));data,located=locate(data,manifest);assert located['summary']['project_records']==2
    report=audit(data,manifest);assert report['status']=='passed';results['source_to_data_full_coverage']=True
    missing=copy.deepcopy(data);missing['people'][0]['fields']['projects'][0]['resp']='';assert audit(missing,manifest)['issues'];results['missing_before_json_detected']=True
    swapped=copy.deepcopy(data);p=swapped['people'][0]['fields']['projects'];p[0]['resp'],p[1]['resp']=p[1]['resp'],p[0]['resp'];assert audit(swapped,manifest)['issues'];results['cross_project_swap_detected']=True
    fake=copy.deepcopy(data);fake['people'][0]['evidence'][0]['source_quotes'][0]['quote']='虚构原文';denied(lambda:audit(fake,manifest));results['fake_quote_rejected']=True
    candidate={'person_id':'p1','index':0,'client':'示例证券','quote':'为示例证券建设行情平台','source_id':'s1','unit_id':data['people'][0]['record_sources']['projects[1]']['anchor_unit_id']}
    denied(lambda:apply(data,[candidate],manifest));results['cross_project_customer_rejected']=True
    for level,expected in [('本科','学士'),('硕士研究生','硕士'),('博士研究生','博士'),('大专','')]:
        d={'people':[{'id':'x','fields':{'education':[{'level':level,'end':'2020','status':'completed','degree':''},{'level':'博士在读','end':'至今','status':'in_progress','degree':''}]}}]}
        out,_,_=apply(confirmed_fixture(d));assert out['people'][0]['fields']['meta']=={'level':level,'degree':expected};assert out['people'][0]['fields']['education'][1]['degree']==''
    d={'people':[{'id':'x','fields':{'education':[{'level':'博士','end':'2020','status':'completed'},{'level':'本科','end':'2022','status':'completed'}]}}]};assert apply(confirmed_fixture(d))[0]['people'][0]['fields']['meta']['level']=='本科';results['latest_completed_degrees_and_in_progress']=True
    d={'people':[{'id':'x','fields':{'education':[{'level':'本科','end':'2020','status':'completed','degree':'无'}]}}]};assert apply(confirmed_fixture(d))[0]['people'][0]['fields']['meta']['degree']=='无';results['explicit_no_bachelor_degree_preserved']=True
    typed,fr=classify(data,manifest);assert typed['people'][0]['field_facts']['education[0].degree']['origin']=='user_default';results['typed_default_not_direct']=True
    data=typed;report=audit(data,manifest)
    for variant in (0,1):
        t=root/f'template-{variant}.docx';doc=Document();doc.add_paragraph('人员简历');table=doc.add_table(rows=3,cols=2);table.style='Table Grid';doc.save(t)
        mp={'page_break_between_people':False,'paragraphs':[{'index':0,'keys':['meta.name','meta.level','meta.degree']}],
            'tables':[{'index':0,'repeat':[{'key':'education','rows':[0,0],'fields':[{'row':0,'cell':variant,'keys':['school','level','end','degree']}]},{'key':'projects','rows':[1,2],'fields':[{'row':0,'cell':variant,'keys':['name','desc','client']},{'row':1,'cell':variant,'key':'resp'}]}]}]}
        mp['widow_control']=True
        mp['tables'][0]['repeat'][-1]['keep_with_next_rows']=[0]
        mp['tables'][0]['repeat'][-1]['cant_split_rows']=[0]
        if variant==1:
            repeat=mp['tables'][0]['repeat'][-1];repeat.update(rows=[1,1],fields=[{'row':0,'cell':1,'keys':['name','desc','client']}],trailing_rows=[{'rows':[2,2],'when_nonempty':'resp','fields':[{'row':0,'cell':1,'key':'resp','prefix':'职责：'}]}])
        out=root/f'out-{variant}.docx';br=build(t,mp,data,manifest,out,report);assert verify(t,mp,data,manifest,out,report)['person_content_verified']
        with ZipFile(out) as package:
            xml=package.read('word/document.xml');assert b'keepNext' in xml and b'cantSplit' in xml and b'widowControl' in xml

        conflict_data=copy.deepcopy(data);unit=units_for(manifest['sources'][0])[0];conflict_data['source_conflicts']=[{'id':'review-1','reason':'Synthetic unresolved source conflict','source_quotes':[{'source_id':'s1','unit_id':unit['unit_id'],'quote':unit['text']}]}];ca=audit(conflict_data,manifest)
        denied(lambda:build(t,mp,conflict_data,manifest,root/f'conflict-final-{variant}.docx',ca))
        conflict_data['output_mode']='candidate';ca=audit(conflict_data,manifest);assert build(t,mp,conflict_data,manifest,root/f'conflict-candidate-{variant}.docx',ca)['source_audit_status']=='needs_review'
        denied(lambda:build(t,mp,missing,manifest,root/f'stale-{variant}.docx',report))
        bad=copy.deepcopy(data);bad['people'][0]['fields']['unmapped']='不得丢失';ba=audit(bad,manifest);denied(lambda:build(t,mp,bad,manifest,root/f'unmapped-{variant}.docx',ba))
        dump(root/'sources.json',manifest);dump(root/'data.json',data);dump(root/'audit.json',report);dump(root/'map.json',mp);dump(root/'build.json',br);dump(root/'facts.json',fr)
        pdf=root/f'render-{variant}.pdf';pdf.write_bytes(b'fixture used only to test hash invalidation, not a rendered document')
        dump(root/'render.json',{'input_unchanged':True,'docx_sha256':sha(out),'pdf':str(pdf.resolve()),'pages':[]})
        files={k:str((root/v).resolve()) for k,v in {'sources':'sources.json','data':'data.json','source_audit':'audit.json','map':'map.json','build':'build.json','facts':'facts.json','render':'render.json'}.items()};files.update(template=str(t.resolve()),output=str(out.resolve()))
        receipt=create(files);assert receipt['status']=='candidate';denied(lambda:check_receipt(receipt));assert check_receipt(receipt,False)['status']=='passed'
        forged=copy.deepcopy(receipt);forged.update(status='closed',pending=[]);denied(lambda:check_receipt(forged,False))
        saved=(root/'map.json').read_bytes();(root/'map.json').write_text('{}');denied(lambda:check_receipt(receipt,False));(root/'map.json').write_bytes(saved)
        original=source.read_bytes();source.write_text('changed');denied(lambda:verify(t,mp,data,manifest,out,report));source.write_bytes(original)
    results.update(two_actual_template_maps=True,stale_audit_rejected=True,unmapped_business_rejected=True,stale_review_receipt_rejected=True,unreviewed_candidate_not_closed=True,source_mutation_rejected=True,unresolved_conflict_blocks_final=True,explicit_conflict_candidate_remains_unreviewed=True,forged_closed_receipt_rejected=True)
    dump(root/'results.json',{'status':'passed','checks':results});print(json.dumps(results,ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--workdir',type=Path,required=True);run(parser.parse_args().workdir)
