"""Behavioral cache, exact-text, blank-review and sample-isolation regressions."""
import copy
import json
import sys
import tempfile
from argparse import Namespace
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from batch_state import Task, file_hash, read, write
from batch_pipeline import prepare, run, validate_visual
from review_bundle import hydrate
from source_audit import units_for, audit
from template_docx import build, verify
from review_ledger import create
from docx import Document


def refused(call):
    try:
        call()
    except (ValueError, KeyError, IndexError):
        return True
    raise AssertionError('Invalid operation was accepted.')


def main():
    checks = {}
    with tempfile.TemporaryDirectory(prefix='resume-batch-test-') as temporary:
        root = Path(temporary)
        source = root/'original.txt'
        source.write_text('示例甲\n项目甲\n项目介绍：'+('完整业务描述；'*300)+'\n职责：保留成果100%\n项目乙\n项目介绍：乙项目介绍\n职责：乙项目职责\n', encoding='utf-8')
        other = root/'other.txt'
        other.write_text('示例乙\n项目丙\n项目介绍：丙项目介绍\n职责：丙项目职责\n', encoding='utf-8')
        template = root/'template.docx'
        doc = Document();doc.add_paragraph('姓名：');table = doc.add_table(rows=2, cols=1);table.style = 'Table Grid';doc.save(template)
        taskdir = root/'task'
        args = Namespace(input=[source, other], template=template, workdir=taskdir, pdf_order_map=None)
        prepared = prepare(args)
        manifest = read(prepared['sources'])
        by_path = {s['source']: s for s in manifest['sources']}
        first = by_path[str(source.resolve())];second = by_path[str(other.resolve())]
        one, two = units_for(first), units_for(second)
        sid, sid2 = first['id'], second['id']
        def record(sid, values, start, stop):
            return {'source_id': sid, 'anchor_unit_id': values[start]['unit_id'], 'from_unit': values[start]['unit_id'],
                    'through_unit': values[stop]['unit_id'], 'reason': 'Synthetic fixture explicitly defines these project boundaries.'}
        def ref(sid, unit):
            return {'source_id': sid, 'unit_ids': [unit['unit_id']], 'strip_prefix': unit['text'].split('：')[0]+'：'}
        p = {'id': 'p1', 'source_ids': [sid], 'fields': {'meta': {'name': '示例甲', 'sex': ''},
             'projects': [{'name': '项目甲', 'desc': '', 'resp': ''}, {'name': '项目乙', 'desc': '', 'resp': ''}]},
             'record_refs': {'projects[0]': record(sid, one, 1, 3), 'projects[1]': record(sid, one, 4, 6)},
             'field_text_refs': {'projects[0].desc': ref(sid, one[2]), 'projects[0].resp': ref(sid, one[3]),
                                 'projects[1].desc': ref(sid, one[5]), 'projects[1].resp': ref(sid, one[6])},
             'evidence_refs': {'meta.name': [{'source_id': sid, 'unit_id': one[0]['unit_id']}]},
             'blank_review_groups': [{'fields': ['meta.sex'], 'classification': 'source_not_provided',
                                      'reviewer': 'synthetic fixture', 'reason': 'Fixture supplies no sex or identity number.'}]}
        p2 = {'id': 'p2', 'source_ids': [sid2], 'fields': {'meta': {'name': '示例乙', 'sex': ''}, 'projects': [{'name': '项目丙', 'desc': '', 'resp': ''}]},
              'record_refs': {'projects[0]': record(sid2, two, 1, 3)},
              'field_text_refs': {'projects[0].desc': ref(sid2, two[2]), 'projects[0].resp': ref(sid2, two[3])},
              'evidence_refs': {'meta.name': [{'source_id': sid2, 'unit_id': two[0]['unit_id']}]},
              'blank_review_groups': copy.deepcopy(p['blank_review_groups'])}
        draft = {'processing_date': '2026-10-20', 'output_mode': 'candidate', 'source_snapshot': prepared['source_snapshot'], 'people': [p, p2]}
        expanded = hydrate(draft, manifest)
        assert expanded['people'][0]['fields']['projects'][0]['desc'] == '完整业务描述；'*300
        assert expanded['people'][0]['fields']['projects'][0]['resp'] == '保留成果100%'
        assert draft['people'][0]['fields']['projects'][0]['resp'] == ''
        checks['exact_long_text_is_materialized_without_truncation_or_mutating_draft'] = True
        assert audit(expanded, manifest)['status'] == 'passed'
        checks['exact_refs_pass_existing_original_to_field_audit'] = True
        content = Path(next(x['bundle'] for x in read(prepared['index'])['sources'] if x['source_id'] == sid)).read_text('utf-8')
        assert all(u['text'] in content for u in one)
        checks['reading_bundle_contains_every_full_original_unit'] = True
        assert expanded['people'][0]['blank_reviews']['meta.sex']['classification'] == 'source_not_provided'
        checks['explicit_grouped_blank_reviews_expand_only_selected_empty_fields'] = True
        assert expanded['people'][0]['field_facts']['meta.sex']['origin'] == 'not_provided'
        checks['one_explicit_blank_review_updates_coverage_and_typed_facts'] = True
        stale = copy.deepcopy(draft);stale['source_snapshot'][sid]['detail_sha256'] = 'old-extraction'
        refused(lambda: hydrate(stale, manifest));checks['compact_refs_and_blank_reviews_are_bound_to_exact_source_snapshot'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['field_text_refs']['projects[0].resp'] = ref(sid2, two[3])
        refused(lambda: hydrate(bad, manifest));checks['cross_person_text_reference_rejected'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['field_text_refs']['projects[0].resp'] = ref(sid, one[6])
        refused(lambda: hydrate(bad, manifest));checks['cross_project_text_reference_rejected'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['fields']['projects'][0]['resp'] = 'already provided'
        refused(lambda: hydrate(bad, manifest));checks['existing_text_is_never_overwritten_by_reference'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['evidence_refs']['meta.name'][0]['quote'] = 'fabricated'
        refused(lambda: hydrate(bad, manifest));checks['false_exact_quote_rejected'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['blank_review_groups'][0]['fields'] = ['meta.name']
        refused(lambda: hydrate(bad, manifest));checks['nonempty_fact_cannot_be_closed_as_blank'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['field_text_refs']['projects[0].resp']['strip_prefix'] = 'missing label'
        refused(lambda: hydrate(bad, manifest));checks['nonexistent_text_label_removal_rejected'] = True
        bad = copy.deepcopy(draft);bad['people'].append({**copy.deepcopy(p), 'id': 'same-file-other-person'})
        refused(lambda: hydrate(bad, manifest));checks['shared_multi_person_source_requires_explicit_person_scope'] = True
        bad = copy.deepcopy(expanded);bad['people'][0]['fields']['projects'][0]['resp'] = ''
        assert audit(bad, manifest)['status'] != 'passed'
        checks['source_omission_still_fails_original_coverage'] = True
        bad = copy.deepcopy(draft);bad['people'][0]['record_refs']['projects[0]']['through_unit'] = one[6]['unit_id']
        refused(lambda: audit(hydrate(bad, manifest), manifest));checks['scope_crossing_another_project_anchor_rejected'] = True
        mapping = {'page_break_between_people': True, 'paragraphs': [{'index': 0, 'key': 'meta.name', 'prefix': '姓名：'}],
                   'tables': [{'index': 0, 'fields': [{'row': 0, 'cell': 0, 'key': 'meta.sex'}],
                               'repeat': [{'key': 'projects', 'rows': [1, 1], 'fields': [{'row': 0, 'cell': 0, 'keys': ['name', 'desc', 'resp']}]}]}]}
        write(root/'draft.json', draft);write(root/'map.json', mapping)
        run_args = Namespace(workdir=taskdir, data=root/'draft.json', map=root/'map.json', clients=None, decisions=None,
                             pilot_review=None, visual=None, previous_receipt=None, soffice=None, defer_render=True)
        converted = run(run_args)
        assert converted['status'] == 'layout_sample_built_not_rendered'
        checks['unified_pipeline_runs_actual_source_gates_before_layout_sample'] = True
        core_data = read(converted['data']);core_audit = read(converted['source-audit'])
        sample_build = read(converted['sample']['build'])
        assert sample_build['sample_only'] and sample_build['person_content_verified']
        checks['sample_is_explicitly_distinguished_from_full_delivery'] = True
        subset = root/'one-person.docx'
        subset_report = build(template, mapping, core_data, manifest, subset, core_audit, ['p2'])
        assert subset_report['people'] == 1 and verify(template, mapping, core_data, manifest, subset, core_audit, ['p2'])['sample_only']
        checks['sample_verifies_selected_content_using_full_source_gates'] = True
        files = {'source_audit': converted['source-audit'], 'facts': converted['facts'], 'render': str(root/'dummy-render.json'), 'build': converted['sample']['build']}
        write(root/'dummy-render.json', {})
        refused(lambda: create(files));checks['layout_sample_cannot_close_full_batch_receipt'] = True
        calls_before = sum(e['kind'] == 'program_start' for e in Task(taskdir).state['events'])
        assert run(run_args)['status'] == 'layout_sample_built_not_rendered'
        assert sum(e['kind'] == 'program_start' for e in Task(taskdir).state['events']) == calls_before
        checks['unchanged_review_and_sample_are_reused_after_current_hash_validation'] = True
        cached = Path(converted['sample']['output']);cached.write_bytes(b'tampered cached document')
        rerun = run(run_args)
        assert Path(rerun['sample']['output']).read_bytes()[:2] == b'PK'
        checks['tampered_cached_output_is_not_reused'] = True
        bad = copy.deepcopy(draft)
        for person in bad['people']:person.pop('blank_review_groups')
        write(root/'unreviewed.json', bad);pending = run(Namespace(**{**vars(run_args), 'data': root/'unreviewed.json'}))
        assert pending['status'] == 'awaiting_blank_source_review' and pending['unreviewed_blank_slots'] == 2
        checks['unreviewed_target_blanks_stop_before_rendering'] = True
        missed = copy.deepcopy(draft)
        missed['people'][0]['blank_review_groups'][0]['classification'] = 'available_but_missed'
        write(root/'missed.json', missed)
        assert run(Namespace(**{**vars(run_args), 'data': root/'missed.json'}))['status'] == 'available_information_still_unfilled'
        checks['known_available_information_cannot_be_closed_while_still_blank'] = True
        sources_before = {s['source']: s['id'] for s in read(prepare(args)['sources'])['sources']}
        task = Task(taskdir);marker = len(task.state['events'])
        source.write_text(source.read_text('utf-8')+'新增原文事实必须回查\n', encoding='utf-8')
        revised = prepare(args)
        events = Task(taskdir).state['events'][marker:]
        assert sum(e['kind'] == 'program_start' and e.get('stage', '').startswith('extract-') for e in events) == 1
        assert revised['source_changes']['changed'] == [sid]
        checks['one_changed_original_only_reextracts_that_source'] = True
        assert revised['review_plan']['source_review_required'] == ['p1']
        checks['incremental_review_plan_identifies_people_affected_by_source_change'] = True
        assert {s['source']: s['id'] for s in read(revised['sources'])['sources']} == sources_before
        checks['source_ids_stay_stable_after_content_change'] = True
        refused(lambda: run(run_args));checks['new_original_fact_cannot_pass_with_old_semantic_data'] = True
        pdf = root/'hash-only.pdf';pdf.write_bytes(b'not a real PDF; hash-bound negative fixture only')
        image = root/'hash-only.png';image.write_bytes(b'not a real rendered page')
        render = {'pdf': str(pdf), 'pages': [{'page': 1, 'png': str(image)}]}
        visual = root/'visual.json';write(visual, {'status': 'passed', 'pdf_sha256': file_hash(pdf), 'reviewed_pages': []})
        refused(lambda: validate_visual(visual, render));checks['incomplete_visual_review_cannot_complete_batch'] = True
        timing = Task(taskdir).timing()
        assert timing['wall_seconds'] >= timing['program_stage_seconds'] and timing['cache_hits'] > 0
        checks['observed_timing_and_cache_events_are_recorded_without_claiming_model_time'] = True
    print(json.dumps({'status': 'passed', 'checks': checks}))


if __name__ == '__main__':
    main()
