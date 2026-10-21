import argparse,contextlib,json,sys,tempfile
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,Border,Side
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from template_xlsx import build
from source_audit import audit,dump
from extract_sources import sha
from xlsx_portable import write_plan


def run(root=None):
    checks={}
    if root is not None:Path(root).mkdir(parents=True,exist_ok=False)
    with (tempfile.TemporaryDirectory() if root is None else contextlib.nullcontext(root)) as folder:
        w=Path(folder);source=w/'source.txt';source.write_text('示例\n甲项目\n乙项目',encoding='utf8')
        detail=w/'detail.json';dump(detail,{'units':[{'unit_id':str(i),'text':s,'kind':'text_line','locator':str(i)} for i,s in enumerate(source.read_text('utf8').splitlines())]})
        manifest={'sources':[{'id':'s','source':str(source),'sha256':sha(source),'detail_file':str(detail),'status':'ok'}]}
        data={'processing_date':'2026-10-20','people':[{'id':'p','source_ids':['s'],'fields':{'meta':{'name':'示例'},'projects':[{'name':'甲项目'},{'name':'乙项目'}]},'record_sources':{f'projects[{i}]':{'source_id':'s','anchor_unit_id':str(i+1),'unit_ids':[str(i+1)]} for i in range(2)}}]}
        wb=Workbook();ws=wb.active;ws.title='简历';ws.column_dimensions['B'].width=24;ws.row_dimensions[10].height=31
        style=Font(name='宋体',size=11,bold=True);border=Border(bottom=Side(style='thin',color='000000'))
        for row in ws['B10:F11']:
            for cell in row:cell.font=style;cell.border=border
        ws.merge_cells('B10:C10');ws['B10']='样例';wb.create_sheet('保留')['A1']='原模板其它信息'
        template=w/'template.xlsx';wb.save(template);before=sha(template)
        mapping={'sheet':'简历','fields':[{'cell':'B3','key':'meta.name'}],'static_cells':[{'cell':'E3','text':'=1+1'},{'cell':'E4','text':0}],'clear_ranges':['B3:F13'],'print_area':'B3:F13','repeat':[{'key':'projects','start_row':10,'height':2,'template_range':'B10:F11','fields':[{'row':0,'column':'B','key':'name'}]}]}
        report=build(template,data,mapping,manifest,audit(data,manifest),w/'out.xlsx',None,None,w/'work',backend='openpyxl')
        out=load_workbook(w/'out.xlsx');s=out['简历']
        assert s['B3'].value=='示例' and s['B10'].value=='甲项目' and s['B12'].value=='乙项目'
        checks['source_gated_dynamic_xlsx_generation_without_node_or_host_api']=True
        assert 'B10:C10' in s.merged_cells and 'B12:C12' in s.merged_cells and s['B12'].font.name=='宋体' and s['B12'].font.sz==11 and s['B12'].border.bottom.style=='thin'
        assert s.column_dimensions['B'].width==24 and s.row_dimensions[12].height==31
        checks['prototype_merges_fonts_borders_column_width_and_row_height_are_copied']=True
        assert s['E3'].value=='=1+1' and s['E3'].data_type=='s' and s['E4'].value=='0';checks['formula_like_resume_text_stays_literal_and_zero_is_not_erased']=True
        assert out['保留']['A1'].value=='原模板其它信息' and sha(template)==before
        checks['other_sheets_and_original_template_preserved']=True
        assert 'B$3:$F$13' in str(s.print_area) and report['status']=='built_requires_native_render'
        checks['print_area_written_and_render_not_falsely_claimed']=True
        plan=json.loads((w/'work/plan.json').read_text('utf8'));bad=w/'native-objects.xlsx'
        with ZipFile(template) as a,ZipFile(bad,'w',ZIP_DEFLATED) as b:
            for n in a.namelist():b.writestr(n,a.read(n))
            b.writestr('xl/media/image1.png',b'unsupported-native-object')
        try:write_plan({**plan,'template':str(bad),'output':str(w/'bad.xlsx')})
        except ValueError:pass
        else:raise AssertionError('Unsupported objects silently lost.')
        assert not (w/'bad.xlsx').exists();checks['unsupported_native_objects_rejected_before_export']=True
    print(json.dumps({'status':'passed','checks':checks}))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--workdir',type=Path)
    run(p.parse_args().workdir)
