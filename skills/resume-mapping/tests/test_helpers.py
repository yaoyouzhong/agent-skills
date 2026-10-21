"""Portable form discovery, anchored media, review isolation and visual reuse negatives."""
import copy
import json
import sys
import tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from form_grid import FormGrid,word_table_grid
from media_evidence import extract,compare,attach_reviewed
from review_queue import export,person_gates
from layout_review import plan,complete
from source_audit import dump
from extract_sources import sha


def rejected(call):
    try:call()
    except ValueError:return
    raise AssertionError('Unsafe operation was accepted.')


def run():
    checks={}
    cells=[{'cell':c,'value':v} for c,v in [('A1','姓 名'),('C1','测试人员'),('A3','教育背景'),('A6','开始时间'),('B6','学校'),('A7','2020'),('B7','示例大学')]]
    g=FormGrid({'cells':cells,'merged_ranges':['A1:B1']})
    assert g.at('姓名')[1]['value']=='测试人员'
    assert g.table_records((3,9),{'start':['开始时间'],'school':['学校','毕业院校']},required=['school'])[0][1]=={'start':'2020','school':'示例大学'}
    checks['merged_neighbor_and_nonadjacent_headers']=True
    bad=copy.deepcopy(cells)+[{'cell':'A10','value':'姓名'},{'cell':'B10','value':'另一人员'}]
    rejected(lambda:FormGrid({'cells':bad,'merged_ranges':['A1:B1']}).at('姓名'))
    checks['ambiguous_labels_require_scope']=True
    with tempfile.TemporaryDirectory() as tmp:
        from docx import Document
        from PIL import Image
        root=Path(tmp);img=root/'image.png';Image.new('RGB',(100,60),'white').save(img)
        d=Document();t=d.add_table(rows=1,cols=3);t.cell(0,0).merge(t.cell(0,1)).text='姓名';t.cell(0,2).text='测试人员';d.add_picture(str(img));doc=root/'source.docx';d.save(doc)
        assert word_table_grid(t).at('姓名')[1]['value']=='测试人员'
        checks['word_merged_cells_are_not_duplicated']=True
        from extract_sources import read_docx
        from record_evidence import locate
        # Last project in a Word multi-person document must stop at its own table row.
        rows=Document();pt=rows.add_table(rows=3,cols=5)
        for c,v in zip(pt.rows[0].cells,['序号','项目名称','客户名称','项目起止时间','担任的角色']):c.text=v
        for ri,name in [(1,'甲项目'),(2,'乙项目')]:
            for c,v in zip(pt.rows[ri].cells,[str(ri),name,'客户甲','2020~2021','开发']):c.text=v
        other=rows.add_table(rows=1,cols=2);other.cell(0,0).text='另一人员';other.cell(0,1).text='无关客户'
        src=root/'multi.docx';rows.save(src);detail,text,warnings=read_docx(src)
        dump(root/'multi-detail.json',detail);(root/'multi.txt').write_text(text,encoding='utf8')
        mm={'sources':[{'id':'s','source':str(src),'sha256':sha(src),'text_file':str(root/'multi.txt'),'detail_file':str(root/'multi-detail.json')}]}
        dd={'people':[{'id':'p','source_ids':['s'],'fields':{'projects':[{'name':'甲项目'},{'name':'乙项目'}]}}]}
        located,_=locate(dd,mm)
        from source_audit import units_for
        all_units={u['unit_id']:u for u in units_for(mm['sources'][0])}
        for i in range(2):
            scope=located['people'][0]['record_sources'][f'projects[{i}]']['unit_ids']
            assert len(scope)==5 and all('/row['+str(i+1)+']/' in all_units[u]['locator'] for u in scope)
        checks['word_last_project_cannot_absorb_next_person']=True
        media=extract(doc,root/'media');assert len(media['assets'])==1 and media['assets'][0]['references']
        assert media['assets'][0]['sha256']==sha(img) and sha(doc)==media['source_sha256']
        checks['anchored_office_media_and_original_unchanged']=True
        m={'sources':[{'id':'s','source':str(doc),'sha256':sha(doc),'detail_file':str(root/'detail.json')}]};dump(root/'detail.json',{'units':[]})
        a=media['assets'][0]
        reviewed=[{'asset_id':a['id'],'asset_sha256':a['sha256'],'status':'visually_reviewed','reviewer':'synthetic_fixture',
                   'owner_name':'测试甲','text':'示例大学 测试甲 已毕业','reason':'Private-data-free attachment behavior fixture.'}]
        attached=attach_reviewed(media,m,reviewed,root/'attached')
        assert len(__import__('source_audit').units_for(attached['sources'][0]))==1
        wrong=copy.deepcopy(reviewed);wrong[0]['asset_sha256']='bad'
        rejected(lambda:attach_reviewed(media,m,wrong,root/'attach-reject'))
        checks['reviewed_media_units_bind_owner_pixels_and_original_package']=True
        from field_facts import classify
        media_unit=__import__('source_audit').units_for(attached['sources'][0])[0]
        wrong_person={'processing_date':'2026-10-20','people':[{'id':'x','source_ids':['s'],'fields':{'meta':{'name':'测试乙'},'degree':'学士'},
          'evidence':[{'field':'degree','source_quotes':[{'source_id':'s','unit_id':media_unit['unit_id'],'quote':media_unit['text']}]}]}]}
        rejected(lambda:classify(wrong_person,attached))
        rejected(lambda:__import__('source_audit').audit(wrong_person,attached))
        checks['all_media_fact_fields_reject_another_person_identity']=True
        data={'people':[{'id':'p1','source_ids':['s'],'fields':{'meta':{'name':'测试甲'}}},{'id':'p2','source_ids':['s'],'fields':{'meta':{'name':'测试乙'}}}]}
        clean={'status':'passed','issues':[]};issue={'status':'needs_review','issues':[{'person_id':'p1','field':'meta.name','reason':'=source disagreement','value':'=CMD()'}]}
        gates=person_gates(data,clean,issue,clean,clean,True)
        assert gates[0]['status']=='candidate' and gates[1]['status']=='closed'
        checks['per_person_gates_are_independent']=True
        conflict={'status':'needs_review','issues':[{'conflict':{'person_id':'p1','source_quotes':[{'source_id':'s','quote':'fixture'}]}}]}
        gates=person_gates(data,conflict,clean,clean,clean,True)
        assert gates[0]['status']=='candidate' and gates[1]['status']=='closed'
        checks['nested_source_conflict_does_not_block_other_people']=True
        q=export(data,m,{'facts':issue},root/'queue');assert q['bindings'] and len(q['rows'])==1
        assert "'=CMD()" in (root/'queue/review.csv').read_text(encoding='utf-8-sig')
        checks['human_queue_is_bound_and_csv_formula_safe']=True
        pdf=root/'old.pdf';pdf.write_bytes(b'visual binding fixture only')
        visual=root/'visual.json';dump(visual,{'status':'passed','reviewed_pages':[1],'pdf_sha256':sha(pdf)})
        receipt=root/'receipt.json';dump(receipt,{'bindings':{'visual':{'path':str(visual),'sha256':sha(visual)},'rendered_pdf':{'path':str(pdf),'sha256':sha(pdf)},'page_images':[{'path':str(img),'sha256':sha(img)}]}})
        render={'pdf':str(pdf),'pages':[{'page':1,'png':str(img)}]};reuse=plan(render,receipt);assert reuse['reused_pages']==[1]
        assert complete(render,reuse,{'status':'passed','pdf_sha256':sha(pdf),'reviewed_pages':[]})['status']=='passed'
        checks['identical_reviewed_pixels_can_be_reused']=True
        fresh=root/'fresh.png';Image.new('RGB',(100,60),'black').save(fresh)
        render['pages'][0]['png']=str(fresh);reuse=plan(render,receipt);assert reuse['needs_visual_review']==[1]
        rejected(lambda:complete(render,reuse,{'status':'passed','pdf_sha256':sha(pdf),'reviewed_pages':[]}))
        checks['changed_pixels_require_new_visual_review']=True
        altered=copy.deepcopy(reuse);altered.update(needs_visual_review=[],reused_pages=[1])
        rejected(lambda:complete(render,altered,{'status':'passed','pdf_sha256':sha(pdf),'reviewed_pages':[]}))
        checks['edited_reuse_plan_is_rejected']=True
    print(json.dumps({'status':'passed','checks':checks},ensure_ascii=False))


if __name__=='__main__':run()
